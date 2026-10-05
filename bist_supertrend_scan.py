import os
import warnings
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime

import numpy as np
import pandas as pd
import requests
import yfinance as yf

warnings.filterwarnings("ignore")

# ============================================================
# ⚙️ TELEGRAM AYARLARI
# ============================================================
TELEGRAM_BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN")
TELEGRAM_CHAT_ID = os.environ.get("TELEGRAM_CHAT_ID")
# ============================================================

telegram_session = requests.Session()


def send_telegram_message(text, parse_mode="HTML"):
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        print("❌ TELEGRAM_BOT_TOKEN veya TELEGRAM_CHAT_ID bulunamadı!")
        return False

    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendMessage"
    max_length = 4000
    text_chunks = [text[i : i + max_length] for i in range(0, len(text), max_length)]

    success = True
    for chunk in text_chunks:
        payload = {
            "chat_id": TELEGRAM_CHAT_ID,
            "text": chunk,
            "parse_mode": parse_mode,
            "disable_web_page_preview": True,
        }
        try:
            r = telegram_session.post(url, data=payload, timeout=30)
            if not r.ok:
                print(f"⚠️ Telegram yanıt hatası: {r.status_code} - {r.text}")
                success = False
        except Exception as e:
            print(f"⚠️ Telegram mesaj gönderme hatası: {e}")
            success = False

    return success


def send_telegram_document(file_path, caption=""):
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        return False

    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendDocument"
    try:
        with open(file_path, "rb") as f:
            files = {"document": f}
            data = {"chat_id": TELEGRAM_CHAT_ID, "caption": caption, "parse_mode": "HTML"}
            r = telegram_session.post(url, files=files, data=data, timeout=120)
            if not r.ok:
                print(f"⚠️ Telegram dosya gönderme yanıtı: {r.status_code} - {r.text}")
            return r.ok
    except Exception as e:
        print(f"⚠️ Telegram dosya gönderme hatası: {e}")
        return False


# ---------------------- HİSSE LİSTESİ ----------------------
BIST_SYMBOLS = [
    "A1CAP", "A1YEN", "AAGYO", "ACSEL", "ADEL", "ADESE", "ADGYO", "AEFES", "AFYON", "AGESA",
    "AGHOL", "AGROT", "AGYO", "AHGAZ", "AHSGY", "AKBNK", "AKCNS", "AKENR", "AKFGY", "AKFIS",
    "AKFYE", "AKGRT", "AKHAN", "AKMGY", "AKSA", "AKSEN", "AKSGY", "AKSUE", "AKYHO", "ALARK",
    "ALBRK", "ALCAR", "ALCTL", "ALFAS", "ALGYO", "ALKA", "ALKIM", "ALKLC", "ALTNY",
    "ALVES", "ANELE", "ANGEN", "ANHYT", "ANSGR", "ARASE", "ARCLK", "ARDYZ", "ARENA", "ARFYE",
    "ARMGD", "ARSAN", "ARTMS", "ARZUM", "ASELS", "ASGYO", "ASTOR", "ASUZU", "ATAGY", "ATAKP",
    "ATATP", "ATATR", "ATEKS", "ATLAS", "ATSYH", "AVGYO", "AVHOL", "AVOD", "AVPGY", "AVTUR",
    "AYCES", "AYDEM", "AYEN", "AYES", "AYGAZ", "AZTEK", "BAGFS", "BAHKM", "BAKAB", "BALAT",
    "BALSU", "BANVT", "BARMA", "BASCM", "BASGZ", "BAYRK", "BEGYO", "BERA", "BESLR", "BESTE",
    "BEYAZ", "BFREN", "BIENY", "BIGCH", "BIGEN", "BIGTK", "BIMAS", "BINBN", "BINHO", "BIOEN",
    "BIZIM", "BJKAS", "BLCYT", "BLUME", "BMSCH", "BMSTL", "BNTAS", "BOBET", "BORLS", "BORSK",
    "BOSSA", "BRISA", "BRKO", "BRKSN", "BRKVY", "BRLSM", "BRMEN", "BRSAN", "BRYAT", "BSOKE",
    "BTCIM", "BUCIM", "BULGS", "BURCE", "BURVA", "BVSAN", "BYDNR", "CANTE", "CASA", "CATES",
    "CCOLA", "CELHA", "CEMAS", "CEMTS", "CEMZY", "CEOEM", "CGCAM", "CIMSA", "CLEBI", "CMBTN",
    "CMENT", "CONSE", "COSMO", "CRDFA", "CRFSA", "CUSAN", "CVKMD", "CWENE", "DAGI", "DAPGM",
    "DARDL", "DCTTR", "DENGE", "DERHL", "DERIM", "DESA", "DESPC", "DEVA", "DGATE", "DGGYO",
    "DGNMO", "DIRIT", "DITAS", "DMRGD", "DMSAS", "DNISI", "DOAS", "DOCO", "DOFER",
    "DOFRB", "DOGUB", "DOHOL", "DOKTA", "DSTKF", "DUNYH", "DURDO", "DURKN", "DYOBY", "DZGYO",
    "EBEBK", "ECILC", "ECOGR", "ECZYT", "EDATA", "EDIP", "EFOR", "EGEEN", "EGEGY", "EGEPO",
    "EGGUB", "EGPRO", "EGSER", "EKGYO", "EKIZ", "EKOS", "EKSUN", "ELITE", "EMKEL", "EMNIS",
    "EMPAE", "ENDAE", "ENERY", "ENJSA", "ENKAI", "ENPRA", "ENSRI", "ENTRA", "EPLAS", "ERBOS",
    "ERCB", "EREGL", "ERSU", "ESCAR", "ESCOM", "ESEN", "ETILR", "ETYAT", "EUHOL", "EUKYO",
    "EUPWR", "EUREN", "EUYO", "EYGYO", "FADE", "FENER", "FLAP", "FMIZP", "FONET", "FORMT",
    "FORTE", "FRIGO", "FRMPL", "FROTO", "FZLGY", "GARAN", "GARFA", "GATEG", "GEDIK", "GEDZA",
    "GENIL", "GENKM", "GENTS", "GEREL", "GESAN", "GIPTA", "GLBMD", "GLCVY", "GLRMK", "GLRYH",
    "GLYHO", "GMTAS", "GOKNR", "GOLTS", "GOODY", "GOZDE", "GRNYO", "GRSEL", "GRTHO", "GSDDE",
    "GSDHO", "GSRAY", "GUBRF", "GUNDG", "GWIND", "GZNMI", "HALKB", "HATEK", "HATSN", "HDFGS",
    "HEDEF", "HEKTS", "HKTM", "HLGYO", "HOROZ", "HRKET", "HTTBT", "HUBVC", "HUNER", "HURGZ",
    "ICBCT", "ICUGS", "IDGYO", "IEYHO", "IHAAS", "IHEVA", "IHGZT", "IHLAS", "IHLGM", "IHYAY",
    "IMASM", "INDES", "INFO", "INGRM", "INTEK", "INTEM", "INVEO", "INVES", "ISATR", "ISBIR",
    "ISBTR", "ISCTR", "ISDMR", "ISFIN", "ISGSY", "ISGYO", "ISKPL", "ISKUR", "ISMEN", "ISSEN",
    "ISYAT", "IZENR", "IZFAS", "IZINV", "IZMDC", "JANTS", "KAPLM", "KAREL", "KARSN", "KARTN",
    "KATMR", "KAYSE", "KBORU", "KCAER", "KCHOL", "KENT", "KERVN", "KFEIN", "KGYO", "KIMMR",
    "KLGYO", "KLKIM", "KLMSN", "KLNMA", "KLRHO", "KLSER", "KLSYN", "KLYPV", "KMPUR", "KNFRT",
    "KOCMT", "KONKA", "KONTR", "KONYA", "KOPOL", "KORDS", "KOTON", "KRDMA", "KRDMB", "KRDMD",
    "KRGYO", "KRONT", "KRPLS", "KRSTL", "KRTEK", "KRVGD", "KSTUR", "KTLEV", "KTSKR", "KUTPO",
    "KUVVA", "KUYAS", "KZBGY", "KZGYO", "LIDER", "LIDFA", "LILAK", "LINK", "LKMNH", "LMKDC",
    "LOGO", "LRSHO", "LUKSK", "LXGYO", "LYDHO", "LYDYE", "MAALT", "MACKO", "MAGEN", "MAKIM",
    "MAKTK", "MANAS", "MARBL", "MARKA", "MARMR", "MARTI", "MAVI", "MCARD", "MEDTR", "MEGAP",
    "MEGMT", "MEKAG", "MEPET", "MERCN", "MERIT", "MERKO", "METRO", "MEYSU", "MGROS", "MHRGY",
    "MIATK", "MMCAS", "MNDRS", "MNDTR", "MOBTL", "MOGAN", "MOPAS", "MPARK", "MRGYO", "MRSHL",
    "MSGYO", "MTRKS", "MTRYO", "MZHLD", "NATEN", "NETAS", "NETCD", "NIBAS", "NTGAZ", "NTHOL",
    "NUGYO", "NUHCM", "OBAMS", "OBASE", "ODAS", "ODINE", "OFSYM", "ONCSM", "ONRYT", "ORCAY",
    "ORGE", "ORMA", "OSMEN", "OSTIM", "OTKAR", "OTTO", "OYAKC", "OYAYO", "OYLUM", "OYYAT",
    "OZATD", "OZGYO", "OZKGY", "OZRDN", "OZSUB", "OZYSR", "PAGYO", "PAHOL", "PAMEL", "PAPIL",
    "PARSN", "PASEU", "PATEK", "PCILT", "PEKGY", "PENGD", "PENTA", "PETKM", "PETUN", "PGSUS",
    "PINSU", "PKART", "PKENT", "PLTUR", "PNLSN", "PNSUT", "POLHO", "POLTK", "PRDGS", "PRKAB",
    "PRKME", "PRZMA", "PSDTC", "PSGYO", "QNBFK", "QNBTR", "QUAGR", "RALYH", "RAYSG", "REEDR",
    "RGYAS", "RNPOL", "RODRG", "RTALB", "RUBNS", "RUZYE", "RYGYO", "RYSAS", "SAFKR", "SAHOL",
    "SAMAT", "SANEL", "SANFM", "SANKO", "SARKY", "SASA", "SAYAS", "SDTTR", "SEGMN", "SEGYO",
    "SEKFK", "SEKUR", "SELEC", "SELVA", "SERNT", "SEYKM", "SILVR", "SISE", "SKBNK", "SKTAS",
    "SKYLP", "SKYMD", "SMART", "SMRTG", "SMRVA", "SNGYO", "SNICA", "SNPAM", "SODSN", "SOKE",
    "SOKM", "SONME", "SRVGY", "SUMAS", "SUNTK", "SURGY", "SUWEN", "SVGYO", "TABGD", "TARKM",
    "TATEN", "TATGD", "TAVHL", "TBORG", "TCELL", "TCKRC", "TDGYO", "TEHOL", "TEKTU", "TERA",
    "TEZOL", "TGSAS", "THYAO", "TKFEN", "TKNSA", "TLMAN", "TMPOL", "TMSN", "TNZTP", "TOASO",
    "TRALT", "TRCAS", "TRENJ", "TRGYO", "TRHOL", "TRILC", "TRMET", "TSGYO", "TSKB", "TSPOR",
    "TTKOM", "TTRAK", "TUCLK", "TUKAS", "TUPRS", "TUREX", "TURGG", "TURSG", "UCAYM", "UFUK",
    "ULAS", "ULKER", "ULUFA", "ULUSE", "ULUUN", "UMPAS", "UNLU", "USAK", "VAKBN", "VAKFA",
    "VAKFN", "VAKKO", "VANGD", "VBTYZ", "VERTU", "VERUS", "VESBE", "VESTL", "VKFYO", "VKGYO",
    "VKING", "VRGYO", "VSNMD", "YAPRK", "YATAS", "YAYLA", "YBTAS", "YEOTK", "YESIL", "YGGYO",
    "YIGIT", "YKBNK", "YKSLN", "YUNSA", "YYAPI", "YYLGD", "ZEDUR", "ZERGY", "ZGYO",
    "ZOREN", "ZRGYO"
]
BIST_SYMBOLS = sorted(list(set(BIST_SYMBOLS)))


def calculate_supertrend(df, period=10, multiplier=3.0):
    high = df["High"]
    low = df["Low"]
    close = df["Close"]

    # ATR Calculation
    tr1 = pd.DataFrame(high - low)
    tr2 = pd.DataFrame(abs(high - close.shift(1)))
    tr3 = pd.DataFrame(abs(low - close.shift(1)))
    tr = pd.concat([tr1, tr2, tr3], axis=1).max(axis=1)
    atr = tr.ewm(alpha=1 / period, adjust=False).mean()

    hl2 = (high + low) / 2
    basic_upper = hl2 + (multiplier * atr)
    basic_lower = hl2 - (multiplier * atr)

    upper_band = pd.Series(0.0, index=df.index)
    lower_band = pd.Series(0.0, index=df.index)
    trend = pd.Series(1, index=df.index)

    for i in range(1, len(df)):
        if basic_upper.iloc[i] < upper_band.iloc[i - 1] or close.iloc[i - 1] > upper_band.iloc[i - 1]:
            upper_band.iloc[i] = basic_upper.iloc[i]
        else:
            upper_band.iloc[i] = upper_band.iloc[i - 1]

        if basic_lower.iloc[i] > lower_band.iloc[i - 1] or close.iloc[i - 1] < lower_band.iloc[i - 1]:
            lower_band.iloc[i] = basic_lower.iloc[i]
        else:
            lower_band.iloc[i] = lower_band.iloc[i - 1]

        if trend.iloc[i - 1] == 1:
            if close.iloc[i] < lower_band.iloc[i]:
                trend.iloc[i] = -1
            else:
                trend.iloc[i] = 1
        else:
            if close.iloc[i] > upper_band.iloc[i]:
                trend.iloc[i] = 1
            else:
                trend.iloc[i] = -1

    supertrend = pd.Series(np.where(trend == 1, lower_band, upper_band), index=df.index)
    return supertrend, trend


def scan_supertrend(symbol):
    try:
        ticker = yf.Ticker(f"{symbol}.IS")
        df = ticker.history(period="6mo", interval="1d")

        if df is not None and len(df) >= 30:
            df = df.dropna(subset=["Close", "High", "Low"])
            st_val, trend = calculate_supertrend(df, period=10, multiplier=3.0)

            df["Supertrend"] = st_val
            df["Trend"] = trend

            last_bar = df.iloc[-1]
            prev_bar = df.iloc[-2]

            # Koşul: Dün Trend == -1 (SAT) iken Bugün Trend == 1 (AL)
            if prev_bar["Trend"] == -1 and last_bar["Trend"] == 1:
                st_stop = last_bar["Supertrend"]
                close_price = last_bar["Close"]
                stop_dist_pct = ((close_price - st_stop) / close_price) * 100

                return {
                    "Hisse": symbol,
                    "Kapanış Fiyatı": round(close_price, 2),
                    "Supertrend Stop": round(st_stop, 2),
                    "Stop Mesafesi (%)": round(stop_dist_pct, 2),
                    "Hacim": int(last_bar["Volume"]),
                }
    except Exception as e:
        print(f"❌ {symbol}: {type(e).__name__}: {e}")
    return None


def build_telegram_message(signals, scan_time_str, total_symbols):
    if not signals:
        return (
            f"📊 <b>BIST Supertrend (10,3) Taraması</b>\n"
            f"⏰ {scan_time_str}\n"
            f"🔎 Taranan: {total_symbols} hisse\n\n"
            f"❌ Yeni Supertrend AL sinyali veren hisse bulunamadı."
        )
    lines = [
        "🟢 <b>BIST Supertrend AL Sinyali Verenler</b>",
        f"⏰ <i>{scan_time_str}</i>",
        f"🔎 Taranan: {total_symbols} hisse",
        f"🎯 <b>Bulunan: {len(signals)} hisse</b>",
        "",
        "━━━━━━━━━━━━━━━━━━━━",
    ]
    for s in signals:
        lines.append(
            f"🚀 <b>{s['Hisse']}</b>\n"
            f"   💰 Fiyat: <code>{s['Kapanış Fiyatı']:.2f}</code> TL\n"
            f"   🛡 Supertrend Destek: <code>{s['Supertrend Stop']:.2f}</code> TL\n"
            f"   📏 Stop Risk Mesafesi: <code>%{s['Stop Mesafesi (%)']:.2f}</code>"
        )
    lines.append("━━━━━━━━━━━━━━━━━━━━")
    lines.append("<i>⚠ Yatırım tavsiyesi değildir.</i>")
    return "\n".join(lines)


if __name__ == "__main__":
    scan_time = datetime.now()
    scan_time_str = scan_time.strftime("%d.%m.%Y %H:%M:%S")

    print("🚀 Supertrend AL Sinyali Taraması Başlatılıyor...")
    print(f"⏰ Tarama Zamanı: {scan_time_str}")
    print("=" * 70)

    signals = []
    processed_count = 0

    with ThreadPoolExecutor(max_workers=5) as executor:
        future_to_symbol = {executor.submit(scan_supertrend, sym): sym for sym in BIST_SYMBOLS}

        for future in as_completed(future_to_symbol):
            processed_count += 1
            res = future.result()
            print(
                f"\rİlerleme: %{int((processed_count / len(BIST_SYMBOLS)) * 100)} ({processed_count}/{len(BIST_SYMBOLS)})",
                end="",
                flush=True,
            )
            if res:
                signals.append(res)

    print("\n\n" + "=" * 70)
    if signals:
        # Stop riski en düşük olandan en yükseğe doğru sırala
        signals = sorted(signals, key=lambda x: x["Stop Mesafesi (%)"])

    # --- EXCEL KAYIT BLOĞU ---
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    filename = f"BIST_Supertrend_Tarama_{scan_time.strftime('%Y%m%d_%H%M')}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Supertrend Taramasi"

    ws["A1"] = f"Tarama Zamanı: {scan_time_str}"
    ws["A1"].font = Font(name="Consolas", size=8, italic=True, bold=True, color="555555")

    headers = ["Hisse", "Kapanış Fiyatı", "Supertrend Stop", "Stop Mesafesi (%)", "Hacim"]
    start_row = 3

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = Font(name="Consolas", size=8, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E4D2B", end_color="1E4D2B", fill_type="solid")  # Koyu Yeşil Header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for row_idx, data in enumerate(signals, start=start_row + 1):
        row_values = [
            data["Hisse"],
            data["Kapanış Fiyatı"],
            data["Supertrend Stop"],
            data["Stop Mesafesi (%)"],
            data["Hacim"],
        ]
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Consolas", size=8)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "right")

            if col_idx in [2, 3]:
                cell.number_format = "#,##0.00"
            elif col_idx == 4:
                cell.number_format = "0.00"
            elif col_idx == 5:
                cell.number_format = "#,##0"

    max_row = start_row + len(signals)
    ws.auto_filter.ref = f"A{start_row}:E{max(max_row, start_row + 1)}"

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col if cell.row >= start_row)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

    wb.save(filename)

    # --- TELEGRAM GÖNDERİM ---
    msg_text = build_telegram_message(signals, scan_time_str, len(BIST_SYMBOLS))
    send_telegram_message(msg_text)

    if signals:
        send_telegram_document(filename, caption=f"📎 BIST Supertrend Tarama - {scan_time_str}")